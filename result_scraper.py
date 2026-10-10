"""
result_scraper.py - marksheet data from the university result site into Excel.

Step 1 (no browser): read the attestation PDFs of a class -> one Excel file with a row per student.
Step 2 (browser):    for every roll number in that Excel file, open the marksheet on the result
                     site and add every field to the same row (one column per field).

All settings live in result_jobs.py.

  python result_scraper.py auto         look for new attestation PDFs in Documents, then do steps 1 + 2
                                        for every class found (main and ATKT), session by session
  python result_scraper.py watch [min]  the same, repeated every N minutes (default WATCH_MINUTES)
  python result_scraper.py              ask for one class (or use CLASSES), run both steps
  python result_scraper.py extract      step 1 only (re-reads the PDFs)
  python result_scraper.py test         step 1 + step 2 for the first 3 roll numbers only
  python result_scraper.py dump FILE    show the text of an attestation PDF and what was understood
  python result_scraper.py report [JUN-2023 [LL336]]   result report + statistics of every class (see result_report.py)

Setup:  pip install pdfplumber openpyxl playwright   (then: playwright install chromium)
"""
import ctypes
import json
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from playwright.sync_api import sync_playwright, Error as PWError, TimeoutError as PWTimeout

import result_jobs as cfg

# ----------------------------------------------------------------------------
# Settings (from result_jobs.py)
# ----------------------------------------------------------------------------
START_URL = cfg.START_URL
CLASSES = list(getattr(cfg, "CLASSES", []))
COURSE_MAP = {k.lower(): v for k, v in getattr(cfg, "COURSE_MAP", {"llb": ("LL", "3")}).items()}
MAIN_STATUSES = [s.upper() for s in getattr(cfg, "MAIN_STATUSES", ["REGULAR", "EXR"])]
ATKT_STATUSES = [s.upper() for s in getattr(cfg, "ATKT_STATUSES", ["ATKT"])]
HEADLESS = getattr(cfg, "HEADLESS", False)
DELAY = getattr(cfg, "DELAY", (0.8, 1.6))
DOCS_ROOT = getattr(cfg, "DOCS_ROOT", None)                         # None = My Documents
RETRY_NOT_FOUND_HOURS = getattr(cfg, "RETRY_NOT_FOUND_HOURS", 24)   # auto: look again at "not found" rolls
RETRY_MISSING_CODE_HOURS = getattr(cfg, "RETRY_MISSING_CODE_HOURS", 6)  # auto: exam not published yet
WATCH_MINUTES = getattr(cfg, "WATCH_MINUTES", 30)
CHANCES = getattr(cfg, "CHANCES", 2)                     # 2 = also look in the "II CHANCE" exams (A2LL324, 2LL324)
CHANCE_FOR_MAIN = getattr(cfg, "CHANCE_FOR_MAIN", True)  # main exams too: LL324, then 2LL324
EXTRA_CODES = {k.upper(): [c.upper() for c in v] for k, v in getattr(cfg, "EXTRA_CODES", {}).items()}
SETTLE_SECONDS = 30      # a PDF written less than this long ago may still be incomplete

# The session being worked on - set by use_session()
SESSION = ""
EXAM_YEAR = ""
SESS = ""
ATT_FOLDER = Path(".")
OUT_FOLDER = Path(".")


def use_session(session, att_folder, out_folder, year=None):
    """JUN-2026 -> exam year 2026, Excel/PDF folders for that session."""
    global SESSION, EXAM_YEAR, SESS, ATT_FOLDER, OUT_FOLDER
    SESSION = session
    SESS = session.replace("-", "_")
    m = re.search(r"\d{4}", session)
    EXAM_YEAR = str(year) if year else (m.group(0) if m else "")
    ATT_FOLDER = Path(att_folder)
    OUT_FOLDER = Path(out_folder)


def session_matches(text):
    """Does an exam name from the result list belong to SESSION? 'JUN-2026' matches 'JUNE-2026'."""
    m = re.fullmatch(r"([A-Za-z]{3})\W*(\d{4})", SESSION.strip())
    if not m:
        return False
    return re.search(rf"{m.group(1).upper()}[A-Z]*[\W_]*{m.group(2)}", text.upper()) is not None


def documents_dir():
    """Windows 'My Documents' folder (also works if it is redirected to OneDrive)."""
    try:
        buf = ctypes.create_unicode_buffer(260)
        ctypes.windll.shell32.SHGetFolderPathW(None, 5, None, 0, buf)  # 5 = My Documents
        return Path(buf.value)
    except Exception:
        return Path.home() / "Documents"


use_session(getattr(cfg, "SESSION", ""), getattr(cfg, "ATT_FOLDER", "."), getattr(cfg, "OUT_FOLDER", "."),
            getattr(cfg, "EXAM_YEAR", None))

MARK_COLS = ["Theory I", "Theory II", "Theory III", "Sessional I", "Sessional II",
             "Practical I", "Practical II", "Total"]
STUDENT_COLS = ["Roll No", "Name", "Father's Name", "Mother's Name", "Enrollment No",
                "Regular/Backlog", "Exam Name", "Exam Centre", "College Name", "Gender",
                "Category", "Address", "Papers", "Class", "Exam Code",
                "Att Regular/Private", "Att Status", "Att Branch", "Source File"]
RESULT_COLS = ["Result Status", "Result Note", "Fetched At", "Result Exam Code"]
DONE_STATUSES = ("OK", "NOT FOUND")
NOT_FOUND_WORDS = ("not found", "no record", "invalid", "does not exist", "no result",
                   "not available", "incorrect", "wrong")
STATE = {"tab_wait": 2500}


class ConfigError(Exception):
    """Something in result_jobs.py or on the page needs a human."""


class CodeMissing(ConfigError):
    """The exam code is not on the result list (results not published yet?)."""


class BrowserLost(Exception):
    pass


# ----------------------------------------------------------------------------
# Which class? ("llb 3rd sem atkt" -> ALL323)
# ----------------------------------------------------------------------------
def parse_class(text):
    """EXCODE = [A for ATKT] + course letters + duration digit + year digit + semester digit,
    e.g. ALL323 = LL.B. (3YDC) second year, third semester, ATKT; LL335 = fifth semester main exam."""
    t = text.strip()
    compact = t.upper().replace(" ", "")
    m = re.fullmatch(r"(A?)(\d?)([A-Z]+)(\d)(\d)(\d)", compact)
    if m and " " not in t:                                   # a code typed directly (A2LL324 = ATKT, II chance)
        atkt, sem = m.group(1) == "A", int(m.group(6))
        return {"code": compact, "sem": sem, "atkt": atkt, "course": None,
                "label": f"{compact} ({'ATKT' if atkt else 'main'} exam, sem {sem})"}

    low = t.lower()
    sem_m = re.search(r"(\d+)\s*(?:st|nd|rd|th)?\s*sem", low) or re.search(r"sem(?:ester)?\s*(\d+)", low)
    if not sem_m:
        raise ConfigError(f"Cannot find the semester in '{text}'. Write e.g. 'llb 3 sem atkt'.")
    sem = int(sem_m.group(1))
    course = next((k for k in COURSE_MAP if re.search(rf"\b{re.escape(k)}\b", low)), None)
    if not course:
        raise ConfigError(f"Unknown course in '{text}'. Known: {list(COURSE_MAP)}. "
                          f"Add it to COURSE_MAP in result_jobs.py, or type the exam code itself.")
    letters, duration = COURSE_MAP[course]
    atkt = "atkt" in low
    year = (sem + 1) // 2
    code = ("A" if atkt else "") + letters + duration + str(year) + str(sem)
    return {"code": code, "sem": sem, "atkt": atkt, "course": course,
            "label": f"{course.upper()} sem {sem} {'ATKT' if atkt else 'main'} ({code})"}


def code_variants(cls):
    """Exam codes to look in for a roll number, in order. The same class is examined again as a
    "II CHANCE" exam with a chance digit in the code:
        ATKT class ALL324 -> ALL324, A2LL324 (ATKT, II chance), 2LL324 (II chance)
        main class  LL324 ->  LL324, 2LL324
    (III chance: set CHANCES = 3 in result_jobs.py). EXTRA_CODES = {"ALL324": ["X1"]} adds odd ones by hand.
    A code typed with a chance digit (A2LL324) is used alone."""
    code = cls["code"]
    core = code[1:] if cls["atkt"] else code
    out = [code]
    if re.fullmatch(r"[A-Z]+\d{3}", core):
        for n in range(2, CHANCES + 1):
            if cls["atkt"]:
                out += [f"A{n}{core}", f"{n}{core}"]
            elif CHANCE_FOR_MAIN:
                out.append(f"{n}{core}")
    out += [c for c in EXTRA_CODES.get(code.upper(), []) if c not in out]
    return out


# ----------------------------------------------------------------------------
# Step 1: attestation PDFs -> students
# ----------------------------------------------------------------------------
ROLL_RE = re.compile(r"Roll\s*No\.?\s*:?\s*(\d{5,})", re.I)
LABEL_RE = re.compile(
    r"(?<![A-Za-z])(Enrollment\s*No\.?|Regular\s*/\s*Backlog|Exam\s*Name|Exam\s*Cent(?:re|er)|College\s*N\w*"
    r"|Father['’]?s\s*Name|Mother['’]?s\s*Name|Gender|Category|Address|Name)\s*:?", re.I)
WS_RE = re.compile(r"\s+")
PAPER_RE = re.compile(r"([^\[\]]+?)\s*\[([A-Za-z]+\d+[A-Za-z0-9]*)\]")


def label_name(raw):
    k = re.sub(r"[^a-z/]", "", raw.lower())
    if k.startswith("enrollment"):
        return "Enrollment No"
    if k.startswith("regular"):
        return "Regular/Backlog"
    if k == "examname":
        return "Exam Name"
    if k.startswith("examcent"):
        return "Exam Centre"
    if k.startswith("college"):
        return "College Name"
    if k.startswith("father"):
        return "Father's Name"
    if k.startswith("mother"):
        return "Mother's Name"
    return {"gender": "Gender", "category": "Category", "address": "Address", "name": "Name"}.get(k, raw)


def pdf_text(path):
    out = []
    with pdfplumber.open(path) as pdf:
        for pg in pdf.pages:
            out.append(pg.extract_text() or "")
    return "\n".join(out)


def parse_attestation(text):
    """One dict per student found in the text of an attestation PDF."""
    marks = list(ROLL_RE.finditer(text))
    students, seen = [], set()
    for i, m in enumerate(marks):
        roll = m.group(1)
        if roll in seen:
            continue
        seen.add(roll)
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        chunk = re.sub(r"\s+", " ", text[m.end():end]).strip()

        parts = re.split(r"Paper\s*Description", chunk, maxsplit=1, flags=re.I)
        head, tail = parts[0], (parts[1] if len(parts) > 1 else "")

        rec = {"Roll No": roll}
        found = list(LABEL_RE.finditer(head))
        for j, lm in enumerate(found):
            nxt = found[j + 1].start() if j + 1 < len(found) else len(head)
            value = head[lm.end():nxt].strip(" :|-")
            rec.setdefault(label_name(lm.group(1)), value)
        if rec.get("Enrollment No"):
            rec["Enrollment No"] = rec["Enrollment No"].split()[0]
        if len(rec.get("Address", "")) > 300:
            rec["Address"] = rec["Address"][:300]

        tail = re.sub(r"(?<![A-Za-z])(Date|Signature|Answer\s*Sheet\s*No\.?|Invigilator['’]?s?\s*(?:Name|Sign))(?![A-Za-z])",
                      " ", tail, flags=re.I)
        papers = [f"{WS_RE.sub(' ', n).strip(' -:.')} [{c}]" for n, c in PAPER_RE.findall(tail)]
        rec["Papers"] = "; ".join(papers)
        students.append(rec)
    return students


def branch_matches(branch, course):
    """True if the PDF's branch belongs to this course. '--All--' (old sessions, all branches in one PDF)
    matches any course when COURSE_MAP has only one course."""
    b = branch.lower()
    if b.startswith(course.lower()):
        return True
    return b.strip("-_ ") == "all" and len(COURSE_MAP) == 1


def select_files(sem, atkt, course=None):
    """Attestation PDFs of this semester whose Status matches the exam type (and whose branch starts
    with the course name, e.g. LLBC254 for llb). File names look like 3SEM_REGULAR_EXR_LLBC254.pdf
    (semester_regular/private_status_branch). Sub-folders are searched too."""
    statuses = ATKT_STATUSES if atkt else MAIN_STATUSES
    chosen, skipped = [], []
    for f in sorted(ATT_FOLDER.rglob("*.pdf")):
        p = f.stem.split("_")
        branch = "_".join(p[3:]).lower()
        if (len(p) >= 3 and p[0].upper() == f"{sem}SEM" and p[2].upper() in statuses
                and (not course or branch_matches(branch, course))):
            chosen.append((f, p))
        else:
            skipped.append(f.name)
    return chosen, skipped


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------
class Book:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            self.wb = load_workbook(self.path)
            self.ws = self.wb["Students"]
            new = False
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Students"
            new = True
        self.cols = {}
        for c in range(1, self.ws.max_column + 1):
            v = self.ws.cell(1, c).value
            if v:
                self.cols[str(v)] = c
        if new:
            for name in STUDENT_COLS + RESULT_COLS:
                self.col(name)
        self.index = {}
        for r in range(2, self.ws.max_row + 1):
            v = self.ws.cell(r, self.cols["Roll No"]).value
            if v is not None:
                self.index[str(v)] = r
        self.raw = self.wb["Raw"] if "Raw" in self.wb.sheetnames else None
        self.raw_index = {}
        if self.raw:
            for r in range(2, self.raw.max_row + 1):
                v = self.raw.cell(r, 1).value
                if v is not None:
                    self.raw_index[str(v)] = r

    def col(self, name):
        if name not in self.cols:
            idx = max(self.cols.values(), default=0) + 1
            self.cols[name] = idx
            cell = self.ws.cell(1, idx, name)
            cell.font = Font(bold=True)
        return self.cols[name]

    def set(self, r, name, value):
        self.ws.cell(r, self.col(name), value)

    def get(self, r, name):
        idx = self.cols.get(name)
        return self.ws.cell(r, idx).value if idx else None

    def has(self, roll):
        return str(roll) in self.index

    def add_student(self, d):
        r = self.ws.max_row + 1
        for k, v in d.items():
            self.set(r, k, v)
        self.index[str(d["Roll No"])] = r
        return r

    def all_rolls(self):
        return sorted(((r, roll) for roll, r in self.index.items()), key=lambda x: x[0])

    def set_raw(self, roll, text):
        if self.raw is None:
            self.raw = self.wb.create_sheet("Raw")
            self.raw.cell(1, 1, "Roll No").font = Font(bold=True)
            self.raw.cell(1, 2, "Marksheet text").font = Font(bold=True)
        r = self.raw_index.get(str(roll)) or self.raw.max_row + 1
        self.raw_index[str(roll)] = r
        self.raw.cell(r, 1, str(roll))
        self.raw.cell(r, 2, text[:30000])

    def save(self):
        self.ws.freeze_panes = "C2"
        while True:
            try:
                self.wb.save(self.path)
                return
            except PermissionError:
                input(f"   '{self.path.name}' is open in Excel. Close it, then press Enter to save... ")


def extract_class(cls, book, state=None):
    files, skipped = select_files(cls["sem"], cls["atkt"], cls.get("course"))
    if not files:
        raise ConfigError(
            f"No attestation PDF in {ATT_FOLDER} matches semester {cls['sem']} / "
            f"{'ATKT' if cls['atkt'] else 'main exam'} (statuses "
            f"{ATKT_STATUSES if cls['atkt'] else MAIN_STATUSES}). Files there: {skipped[:10]}")

    check = state if (state and book.index) else None      # a new/empty Excel needs every PDF
    todo = []
    for f, parts in files:
        if check and check.seen(f):
            continue
        if check and time.time() - f.stat().st_mtime < SETTLE_SECONDS:
            print(f"   (waiting: {f.name} was written a moment ago)")
            continue
        todo.append((f, parts))
    if not todo:
        print(f"   no new attestation PDFs ({len(files)} already in the Excel)")
        return

    print(f"   {len(todo)} attestation PDF(s) to read:")
    added = 0
    for f, parts in todo:
        try:
            studs = parse_attestation(pdf_text(f))
        except Exception as e:                      # broken / half-downloaded / not a real PDF
            print(f"   - {f.name}: SKIPPED, cannot be read as a PDF ({f.stat().st_size} bytes, "
                  f"{type(e).__name__}). Delete it and download it again.")
            continue
        print(f"   - {f.name}: {len(studs)} student(s)")
        if not studs:
            print(f"     WARNING: nothing understood. Run:  python result_scraper.py dump \"{f}\"")
            continue
        for s in studs:
            if book.has(s["Roll No"]):
                continue
            s.update({"Class": cls["label"], "Exam Code": cls["code"],
                      "Att Regular/Private": parts[1], "Att Status": parts[2],
                      "Att Branch": "_".join(parts[3:]), "Source File": f.name})
            book.add_student(s)
            added += 1
        if state:
            state.mark(f)
    book.save()
    print(f"   {added} new student(s) added; {len(book.index)} in total -> {book.path}")


class State:
    """Remembers which PDFs are already in the Excel files, and when an exam code was last missing."""

    def __init__(self, path):
        self.path = Path(path)
        self.data = {"pdfs": {}, "missing": {}}
        if self.path.exists():
            try:
                self.data.update(json.loads(self.path.read_text(encoding="utf-8")))
            except Exception:
                pass

    @staticmethod
    def sig(f):
        s = Path(f).stat()
        return [s.st_size, int(s.st_mtime)]

    def seen(self, f):
        return self.data["pdfs"].get(str(f)) == self.sig(f)

    def mark(self, f):
        self.data["pdfs"][str(f)] = self.sig(f)

    def missing_recently(self, code, hours):
        ts = self.data["missing"].get(code)
        return bool(ts) and (time.time() - ts) < hours * 3600

    def set_missing(self, code):
        self.data["missing"][code] = time.time()

    def clear_missing(self, code):
        self.data["missing"].pop(code, None)

    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.data, indent=1), encoding="utf-8")


# ----------------------------------------------------------------------------
# Step 2: the result site
# ----------------------------------------------------------------------------
def snap(page, tag):
    """Save a screenshot + html of the page, so a failure can be looked at."""
    try:
        d = OUT_FOLDER / "debug"
        d.mkdir(parents=True, exist_ok=True)
        stem = d / f"{datetime.now():%H%M%S}_{tag}"
        page.screenshot(path=str(stem) + ".png")
        Path(str(stem) + ".html").write_text(page.content(), encoding="utf-8")
        print(f"   [debug] saved {stem}.png / .html")
    except Exception:
        pass


def hook_dialogs(ctx, msgs):
    def hook(pg):
        pg.on("dialog", lambda d: (msgs.append(d.message), d.accept()))
    for pg in ctx.pages:
        hook(pg)
    ctx.on("page", hook)


def act_follow(ctx, page, action):
    """Run action() (a click, Enter ...). Returns the page that shows the outcome:
    a new tab if one opened, otherwise the same page."""
    try:
        with ctx.expect_page(timeout=STATE["tab_wait"]) as info:
            action()
        np = info.value
        try:
            np.wait_for_load_state("load", timeout=15000)
        except PWTimeout:
            pass
        return np
    except PWTimeout:
        try:
            page.wait_for_load_state("load", timeout=15000)
        except PWTimeout:
            pass
        return page


def find_roll_input(page):
    """The roll number box: an input with 'roll' in its name/id/placeholder, else the only text box."""
    loc = page.locator("input:not([type]), input[type=text], input[type=number], input[type=tel]")
    cands = []
    for i in range(loc.count()):
        el = loc.nth(i)
        try:
            if not el.is_visible():
                continue
            attrs = " ".join(filter(None, [el.get_attribute("name"), el.get_attribute("id"),
                                           el.get_attribute("placeholder")])).lower()
        except Exception:
            continue
        if "roll" in attrs:
            return el
        if "captcha" in attrs or "search" in attrs:
            continue
        cands.append(el)
    views = page.locator("a:has-text('View Result'), button:has-text('View Result'), "
                         "input[value*='View Result']").count()
    if cands and views <= 1:
        return cands[0]
    return None


def find_submit(page):
    """The button that submits the roll number."""
    btns = page.locator("button:visible, input[type=submit]:visible, input[type=button]:visible")
    for i in range(btns.count()):
        b = btns.nth(i)
        try:
            label = (b.inner_text() or b.get_attribute("value") or "").strip().lower()
        except Exception:
            continue
        if re.search(r"home|back|reset|clear|cancel|close", label):
            continue
        if re.search(r"submit|view|show|search|get|go\b|result|proceed|find|download", label):
            return b
    return None


def captcha_visible(page):
    loc = page.locator("input[name*='captcha' i]:visible, input[id*='captcha' i]:visible, "
                       "input[placeholder*='captcha' i]:visible")
    return loc.count() > 0


def go_to_entry(ctx, page, cls):
    """Start page -> year -> class list -> View Result -> the page with the roll number box."""
    page.goto(START_URL, wait_until="load")

    sel = page.locator("select").first
    try:
        sel.wait_for(state="visible", timeout=20000)
    except PWTimeout:
        snap(page, "no_year_dropdown")
        raise ConfigError("The year dropdown did not appear on the start page.")
    opts = [t.strip() for t in sel.locator("option").all_inner_texts()]
    idx = next((i for i, t in enumerate(opts) if t == EXAM_YEAR), None)
    if idx is None:
        idx = next((i for i, t in enumerate(opts) if t.startswith(EXAM_YEAR)), None)
    if idx is None:
        idx = next((i for i, t in enumerate(opts) if EXAM_YEAR in t), None)
    if idx is None:
        raise ConfigError(f"Year '{EXAM_YEAR}' is not in the Years dropdown: {opts}")
    sel.select_option(index=idx)

    proceed = page.locator("button:has-text('Proceed'), input[value*='Proceed' i], "
                           "a:has-text('Proceed for Result')").first
    page = act_follow(ctx, page, proceed.click)

    # the row of this exam code, and its View Result button
    code = cls["code"]
    view = ("a[contains(normalize-space(.),'View Result')] or .//button[contains(normalize-space(.),'View Result')]"
            " or .//input[contains(@value,'View Result')]")
    rows = page.locator(f"xpath=//*[normalize-space(text())='{code}']/ancestor::*[.//{view}][1]")
    for _ in range(20):
        if rows.count() > 0:
            break
        page.wait_for_timeout(1000)
    if rows.count() == 0:
        snap(page, "code_not_found")
        letters = re.sub(r"\d", "", code)
        letters = letters[1:] if cls["atkt"] else letters
        found = sorted({t.strip() for t in page.locator("td, div").all_inner_texts()
                        if re.fullmatch(r"A?[A-Z]{2,}\d{2,3}", t.strip()) and letters in t})
        raise CodeMissing(f"Exam code {code} is not on the list for {EXAM_YEAR}. "
                          f"Similar codes on the page: {found[:40]}")

    texts = [rows.nth(i).inner_text() for i in range(rows.count())]
    keep = [i for i, t in enumerate(texts) if session_matches(t)]
    pick = keep[0] if keep else 0
    if len(texts) > 1:
        print(f"   {code} appears {len(texts)} times; using: {' '.join(texts[pick].split())[:90]}")
    row = rows.nth(pick)
    button = row.locator("a:has-text('View Result'), button:has-text('View Result'), "
                         "input[value*='View Result']").first
    entry = act_follow(ctx, page, button.click)

    for _ in range(20):
        if find_roll_input(entry):
            return entry
        entry.wait_for_timeout(1000)
    snap(entry, "no_roll_box")
    raise ConfigError("No roll number box found after clicking View Result.")


def ensure_entry(ctx, entry, cls):
    if entry.is_closed():
        raise BrowserLost("entry page closed")
    if find_roll_input(entry):
        return entry
    try:
        entry.go_back(wait_until="load", timeout=10000)
        if find_roll_input(entry):
            return entry
    except Exception:
        pass
    return go_to_entry(ctx, entry, cls)


def wait_marksheet(target, before, dialogs, timeout=25):
    for _ in range(timeout * 2):
        body = ""
        try:
            if target.locator("text=STUDENT DETAILS").count() > 0:
                return "ok", ""
            body = target.inner_text("body")
        except Exception:
            if target.is_closed():
                raise BrowserLost("result page closed")
        if dialogs:
            return "notfound", dialogs[-1]
        if body and body != before:
            low = body.lower()
            for w in NOT_FOUND_WORDS:
                if w in low:
                    line = next((l.strip() for l in body.splitlines() if w in l.lower()), w)
                    return "notfound", line[:150]
        target.wait_for_timeout(500)
    return "timeout", ""


JS_ROWS = """() => [...document.querySelectorAll('tr')]
    .filter(tr => !tr.querySelector('table'))
    .map(tr => [...tr.children].map(td => td.innerText.replace(/\\s+/g, ' ').trim()))"""
ROMAN = {"I", "II", "III", "IV", "V", ""}


def parse_rows(rows, body):
    """Marksheet table rows -> {column: value}."""
    out = {}
    lines = [l.strip() for l in body.splitlines() if l.strip()]
    for i, l in enumerate(lines):
        if "STUDENT DETAILS" in l.upper():
            title = [x for x in lines[max(0, i - 2):i]]
            out["Marksheet Title"] = " | ".join(title)
            break

    mode = None
    for row in rows:
        row = [c for c in row]
        if not any(row):
            continue
        head = row[0].upper()
        if len(row) == 1 or (len(row) > 1 and not any(row[1:])):
            if "STUDENT DETAILS" in head:
                mode = "student"
                continue
            if "MARKS DETAILS" in head:
                mode = "marks"
                continue
            if "FINAL RESULT" in head:
                mode = "final"
                continue
            continue
        if mode == "student" and len(row) >= 2:
            out[f"Marksheet {row[0]}"] = row[1]
        elif mode == "marks":
            if head.startswith("SUBJECT") or all(c.upper() in ROMAN for c in row) \
                    or row[0] in ("Theory", "Sessional", "Practical"):
                continue
            vals = row[1:]
            names = MARK_COLS if len(vals) == len(MARK_COLS) else [f"Col {i + 1}" for i in range(len(vals))]
            for n, v in zip(names, vals):
                if v:
                    out[f"{row[0]} | {n}"] = v
        elif mode == "final" and len(row) >= 2:
            out[f"Final {row[0]}"] = row[1]
    return out


def fetch_one(ctx, entry, roll, dialogs):
    """Returns (status, note, fields, raw_text)."""
    inp = find_roll_input(entry)
    if inp is None:
        raise BrowserLost("roll box vanished")
    inp.fill("")
    inp.fill(str(roll))
    if captcha_visible(entry):
        input("   A captcha is shown. Type it in the browser window, then press Enter here... ")
    before = entry.inner_text("body")
    dialogs.clear()

    btn = find_submit(entry)
    action = btn.click if btn else (lambda: inp.press("Enter"))
    target = act_follow(ctx, entry, action)
    STATE["tab_wait"] = 800 if target is entry else 3000

    try:
        status, note = wait_marksheet(target, before, dialogs)
        if status == "ok":
            rows = target.evaluate(JS_ROWS)
            body = target.inner_text("body")
            fields = parse_rows(rows, body)
            raw = body.split("Disclaimer")[0].strip()
            return "OK", "", fields, raw
        if status == "timeout":
            snap(target, f"timeout_{roll}")
            return "ERROR", "no marksheet appeared within 25 s", {}, ""
        return "NOT FOUND", note, {}, ""
    finally:
        if target is not entry:
            try:
                target.close()
            except Exception:
                pass


def fetch_class(p, cls, book, limit=None, retry_nf_hours=None):
    def due(r):
        status = str(book.get(r, "Result Status") or "")
        if status not in DONE_STATUSES:
            return True
        if status == "NOT FOUND" and retry_nf_hours:          # results may be published later
            try:
                at = datetime.strptime(str(book.get(r, "Fetched At")), "%Y-%m-%d %H:%M:%S")
                return (datetime.now() - at).total_seconds() >= retry_nf_hours * 3600
            except ValueError:
                return True
        return False

    def pending():
        return [(r, roll) for r, roll in book.all_rolls() if due(r)]

    stats = {"OK": 0, "NOT FOUND": 0, "ERROR": 0}
    variants = code_variants(cls)
    todo = pending()
    if limit:
        todo = todo[:limit]
    total = len(todo)
    print(f"   {total} roll number(s) to look up ({len(book.index) - len(pending())} already done)")
    if not total:
        return stats
    if len(variants) > 1:
        print(f"   exam codes tried for each roll, in this order: {', '.join(variants)}")

    tried = {}          # roll -> exam codes it was not found in (kept if the browser has to restart)
    last_note = {}      # roll -> what the site said
    final = set()       # rolls that have their final result in the Excel during this run
    skipped_codes = []  # exam codes that are not on the result list at all
    first_missing = {}
    written = {"n": 0}

    def store(r, roll, status, note, fields, raw, vcode=None):
        for k, v in fields.items():
            book.set(r, k, v)
        book.set(r, "Result Status", status)
        book.set(r, "Result Note", note)
        book.set(r, "Fetched At", f"{datetime.now():%Y-%m-%d %H:%M:%S}")
        if vcode and status == "OK":
            book.set(r, "Result Exam Code", vcode)
        if raw:
            book.set_raw(roll, raw)
        stats[status] += 1
        final.add(roll)
        written["n"] += 1
        if written["n"] % 5 == 0:
            book.save()

    for attempt in range(1, 4):
        todo = [(r, roll) for r, roll in pending() if roll not in final]
        if limit:
            todo = todo[:max(0, limit - sum(stats.values()))]
        if not todo:
            break
        browser = None
        try:
            browser = p.chromium.launch(headless=HEADLESS)
            ctx = browser.new_context(no_viewport=True) if not HEADLESS else browser.new_context()
            dialogs = []
            hook_dialogs(ctx, dialogs)
            page = ctx.new_page()

            for vi, vcode in enumerate(variants):                 # one pass per exam code
                vcls = {**cls, "code": vcode}
                rows_v = [(r, roll) for r, roll in todo if roll not in final and vcode not in tried.get(roll, ())]
                if not rows_v:
                    continue
                if vi > 0:
                    print(f"   --- {len(rows_v)} roll number(s) not found so far: looking in {vcode} ---")
                try:
                    entry = go_to_entry(ctx, page, vcls)
                except CodeMissing as e:
                    if vcode not in skipped_codes:
                        skipped_codes.append(vcode)
                        first_missing.setdefault("msg", str(e))
                        print(f"   {vcode} is not on the result list - nothing to look up there")
                    for _, roll in rows_v:
                        tried.setdefault(roll, set()).add(vcode)
                    continue
                page = entry
                more = vi < len(variants) - 1

                errors_in_row = 0
                for n, (r, roll) in enumerate(rows_v, 1):
                    try:
                        entry = ensure_entry(ctx, entry, vcls)
                        status, note, fields, raw = fetch_one(ctx, entry, roll, dialogs)
                    except BrowserLost:
                        raise
                    except ConfigError:
                        raise
                    except Exception as e:
                        if entry.is_closed():
                            raise BrowserLost(str(e))
                        status, note, fields, raw = "ERROR", f"{type(e).__name__}: {str(e)[:150]}", {}, ""
                        snap(entry, f"error_{roll}")
                        try:
                            entry = go_to_entry(ctx, entry, vcls)
                        except ConfigError:
                            raise
                        except Exception:
                            pass
                    page = entry

                    tag = f"[{vcode}] " if len(variants) > 1 else ""
                    if status == "NOT FOUND":
                        tried.setdefault(roll, set()).add(vcode)
                        last_note[roll] = note
                        errors_in_row = 0
                        print(f"   {tag}[{n}/{len(rows_v)}] {roll}: not found" +
                              ("; will try the next exam code" if more else "") +
                              (f"  ({note})" if note and not more else ""))
                    else:
                        if status == "OK" and vi > 0:
                            chance = re.match(r"A?(\d)", vcode)
                            note = f"found in {vcode}" + (f" (chance {chance.group(1)})" if chance else "")
                        store(r, roll, status, note, fields, raw, vcode)
                        errors_in_row = errors_in_row + 1 if status == "ERROR" else 0
                        extra = (f"{fields.get('Marksheet Name', '')}  {fields.get('Final Total', '')}  "
                                 f"{fields.get('Final Result', '')}") if status == "OK" else note
                        print(f"   {tag}[{n}/{len(rows_v)}] {roll}: {status}  {extra}".rstrip())
                        if errors_in_row >= 5:
                            raise ConfigError("5 errors in a row - stopping. Check the debug files in "
                                              f"{OUT_FOLDER / 'debug'} and tell me what they show.")
                    time.sleep(random.uniform(*DELAY))

            if len(skipped_codes) == len(variants):               # none of the codes is published (yet)
                raise CodeMissing(first_missing.get("msg", f"Exam code {cls['code']} is not on the result list."))

            # every exam code has been tried for the rolls that are still without a result
            checked = [v for v in variants if v not in skipped_codes]
            for r, roll in todo:
                if roll not in final and set(variants) <= tried.get(roll, set()):
                    base = last_note.get(roll) or "not found"
                    store(r, roll, "NOT FOUND", base + (f" [looked in {', '.join(checked)}]" if len(variants) > 1 else ""),
                          {}, "", None)
            break
        except (BrowserLost, PWError) as e:
            msg = (str(e).splitlines() or [''])[0][:120]
            print(f"\n   >>> browser/page problem ({type(e).__name__}: {msg}). Restarting (attempt {attempt}) ...\n")
        finally:
            book.save()
            try:
                if browser:
                    browser.close()
            except Exception:
                pass
    return stats


# ----------------------------------------------------------------------------
# Automatic mode: new attestation PDFs in Documents -> Excel -> marksheets
# ----------------------------------------------------------------------------
def discover_sessions(docs):
    """Session folders made by mponline.py: <MON_YEAR>\\<MON_YEAR>_ATTESTATION  (e.g. JUN_2026)."""
    found = []
    if not docs.is_dir():
        return found
    for d in sorted(docs.iterdir()):
        if d.is_dir() and re.fullmatch(r"[A-Za-z]{3}_\d{4}", d.name):
            att = d / f"{d.name}_ATTESTATION"
            if att.is_dir():
                found.append((d.name, att))
    return found


def classify(att):
    """The classes found in an _ATTESTATION folder: {(course, semester, is_atkt): number of PDFs}.
    Main exam = Regular/EX students, ATKT exam = ATKT students (statuses are set in result_jobs.py)."""
    groups, ignored = {}, {}
    for f in sorted(att.rglob("*.pdf")):
        p = f.stem.split("_")
        m = re.fullmatch(r"(\d+)SEM", p[0].upper()) if len(p) >= 4 else None
        if not m:
            ignored[f.name] = "not a semester file"
            continue
        status, branch = p[2].upper(), "_".join(p[3:])
        course = next((k for k in COURSE_MAP if branch_matches(branch, k)), None)
        if not course:
            ignored[f.name] = f"no course for branch '{branch}' (add it to COURSE_MAP)"
            continue
        if status in ATKT_STATUSES:
            atkt = True
        elif status in MAIN_STATUSES:
            atkt = False
        else:
            ignored[f.name] = f"status {status} is neither a main nor an ATKT status"
            continue
        key = (course, int(m.group(1)), atkt)
        groups[key] = groups.get(key, 0) + 1
    return groups, ignored


def run_auto(p, retry=False):
    docs = Path(DOCS_ROOT) if DOCS_ROOT else documents_dir()
    found = discover_sessions(docs)
    print(f"\n[{datetime.now():%Y-%m-%d %H:%M}] Looking in {docs}")
    if not found:
        print("   No <MON>_<YEAR>\\<MON>_<YEAR>_ATTESTATION folders there yet.")
        return
    print(f"   Sessions found: {', '.join(n.replace('_', '-') for n, _ in found)}")

    for name, att in found:
        use_session(name.replace("_", "-"), att, docs / name / f"{name}_RESULTS")
        groups, ignored = classify(att)
        print(f"\n##### {SESSION}: {len(groups)} class(es) from the attestation PDFs")
        for fname, why in list(ignored.items())[:5]:
            print(f"   ignored {fname}: {why}")
        if len(ignored) > 5:
            print(f"   ... and {len(ignored) - 5} more ignored file(s)")

        state = State(OUT_FOLDER / "state.json")
        for course, sem, atkt in sorted(groups):
            cls = parse_class(f"{course} {sem} sem" + (" atkt" if atkt else ""))
            print(f"\n===== {SESSION}: {cls['label']} =====")
            book = Book(OUT_FOLDER / f"{SESS}_{cls['code']}.xlsx")
            try:
                extract_class(cls, book, state)
                if not retry and state.missing_recently(cls["code"], RETRY_MISSING_CODE_HOURS):
                    print(f"   {cls['code']} was not on the result list recently; will look again later")
                    continue
                stats = fetch_class(p, cls, book, retry_nf_hours=1e-9 if retry else RETRY_NOT_FOUND_HOURS)
                state.clear_missing(cls["code"])
                print(f"   DONE: {stats}  ->  {book.path}")
                if stats.get("OK"):
                    book.save()
                    try:
                        import result_report
                        result_report.build_reports(docs, SESS, cls["code"])
                    except Exception as e:
                        print(f"   (report not updated: {type(e).__name__}: {e})")
            except CodeMissing as e:
                state.set_missing(cls["code"])
                print(f"   Result not published yet ({cls['code']} is not on the list). Will look again later.")
            except ConfigError as e:
                print(f"\n>>> {e}")
            finally:
                book.save()
                state.save()


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def dump(path):
    text = pdf_text(path)
    lines = [l for l in text.splitlines() if l.strip()]
    print(f"--- first 45 text lines of {path} ({len(lines)} lines in total) ---")
    for l in lines[:45]:
        print("  ", l)
    studs = parse_attestation(text)
    print(f"\n--- {len(studs)} student(s) understood; the first one: ---")
    for k, v in (studs[0] if studs else {}).items():
        print(f"   {k}: {v}")


def ask_classes():
    raw = input("Class (e.g. llb 3rd sem atkt, llb 5 sem; or an exam code like ALL323). "
                "Several: separate with commas: ")
    return [x.strip() for x in raw.split(",") if x.strip()]


def main():
    args = sys.argv[1:]
    mode = args[0].lower() if args else "all"
    if mode == "report":
        import result_report
        result_report.build_reports(only_session=args[1] if len(args) > 1 else None,
                                    only_code=args[2] if len(args) > 2 else None)
        return

    if mode == "dump":
        if len(args) < 2:
            sys.exit('Usage: python result_scraper.py dump "C:\\path\\file.pdf"')
        dump(args[1])
        return

    if mode in ("auto", "watch"):
        with sync_playwright() as p:
            if mode == "auto":
                run_auto(p, retry="retry" in [a.lower() for a in args[1:]])
                return
            minutes = float(args[1]) if len(args) > 1 else WATCH_MINUTES
            try:
                while True:
                    try:
                        run_auto(p)
                    except Exception as e:                      # keep watching whatever happens
                        print(f"\n>>> this check failed ({type(e).__name__}: {str(e)[:150]}); will retry")
                    print(f"\nNext check in {minutes:g} minute(s). Press Ctrl+C to stop.")
                    time.sleep(minutes * 60)
            except KeyboardInterrupt:
                print("\nStopped.")
        return

    texts = CLASSES or ask_classes()
    try:
        classes = [parse_class(t) for t in texts]
    except ConfigError as e:
        sys.exit(f"\n>>> {e}")

    with sync_playwright() as p:
        for cls in classes:
            print(f"\n===== {cls['label']} =====")
            book = Book(OUT_FOLDER / f"{SESS}_{cls['code']}.xlsx")
            try:
                if mode == "extract" or not book.index:
                    extract_class(cls, book)
                else:
                    print(f"   Excel already has {len(book.index)} student(s); "
                          f"run 'extract' to read the PDFs again.")
                if mode == "extract":
                    continue
                stats = fetch_class(p, cls, book, limit=3 if mode == "test" else None)
                print(f"   DONE: {stats}  ->  {book.path}")
            except ConfigError as e:
                print(f"\n>>> {e}")
                book.save()


if __name__ == "__main__":
    main()
