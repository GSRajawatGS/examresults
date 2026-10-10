"""
db_sync.py - copy the data you already collect into Supabase tables (see supabase_schema.sql).

  python db_sync.py                       results (attestation + marksheet data) and sitting plans
  python db_sync.py results [JUN-2026 [LL323]]   only the class Excel files (all, one session, one class)
  python db_sync.py sitting [JUN-2026]    only the sitting-plan PDFs
  python db_sync.py attestation [JUN-2026]   only the attestation-sheet PDFs (every course, normal and NEP)
  python db_sync.py counts                how many rows each table holds

Where the data comes from
  results   the class Excel files of result_scraper.py   <Documents>\\JUN_2026\\JUN_2026_RESULTS\\JUN_2026_LL323.xlsx
            -> students (attestation data), marksheets (result per student), marksheet_papers (marks per paper)
  attestation  the attestation PDFs of mponline.py   <Documents>\\JUN_2026\\JUN_2026_ATTESTATION\\...\\2SEM_REGULAR_EXR_LLBC254.pdf
            -> attestation_students (one row per student per PDF: name, father, address, papers ...)
  sitting   the sitting-plan PDFs of mponline.py           <Documents>\\JUN_2026\\JUN_2026_LLB201_...\\2SEM_REGULAR_REGULAR_LLBC254.pdf
            -> sitting_plans (class, paper, roll numbers)
It can be run as often as you like: a row that is already there is updated, never duplicated.

  pip install supabase python-dotenv pdfplumber openpyxl
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import db
import result_report as rr

BASE = Path(__file__).resolve().parent
STATE_FILE = BASE / "db_sync_state.json"

STUDENT_MAP = {
    "Name": "name", "Father's Name": "father_name", "Mother's Name": "mother_name",
    "Enrollment No": "enrollment_no", "Regular/Backlog": "regular_backlog", "Exam Name": "exam_name",
    "Exam Centre": "exam_centre", "College Name": "college_name", "Gender": "gender", "Category": "category",
    "Address": "address", "Papers": "papers", "Class": "class_label",
    "Att Regular/Private": "att_regular_private", "Att Status": "att_status", "Att Branch": "att_branch",
    "Source File": "source_file",
}
MS_MAP = {
    "Result Status": "result_status", "Result Note": "result_note", "Result Exam Code": "result_exam_code",
    "Marksheet Title": "title", "Marksheet Roll Number": "ms_roll_number",
    "Marksheet Enrollment Number": "ms_enrollment", "Marksheet Name": "ms_name",
    "Marksheet Father's Name": "ms_father", "Marksheet Mother's Name": "ms_mother",
    "Marksheet ABCID": "abcid", "Final Result": "final_result",
}
ROLL_TOKEN = re.compile(r"(?<!\d)\d{7,12}(?!\d)")
SESSION_DIR = re.compile(r"([A-Za-z]{3})_(\d{4})(?:_(.+))?")


def clean_val(v):
    if v is None:
        return None
    if isinstance(v, float) and v != v:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def text_val(v):
    v = clean_val(v)
    return None if v is None else str(v)


def to_iso(text):
    """'2026-10-08 12:30:00' (PC time) -> ISO text with the PC's time zone."""
    if not text:
        return None
    try:
        return datetime.strptime(str(text), "%Y-%m-%d %H:%M:%S").astimezone().isoformat()
    except ValueError:
        return None


# ----------------------------------------------------------------------------
# Class Excel file -> rows
# ----------------------------------------------------------------------------
def build_rows(session, code, header, students, raw=None):
    """(student rows, marksheet rows, paper rows) of one class workbook. session = 'JUN-2026'."""
    paper_cols = {}
    for c in header:
        if " | " in c:
            paper, comp = c.rsplit(" | ", 1)
            paper_cols[c] = (paper.strip(), comp.strip())
    mapped = set(MS_MAP) | {"Final Total"}
    stud_rows, ms_rows, pap_rows = [], [], []

    for s in students:
        roll = text_val(s.get("Roll No"))
        if roll is None:
            continue
        key = {"session": session, "exam_code": code, "roll_no": roll}

        row = dict(key)
        for src, dst in STUDENT_MAP.items():
            if src in s:
                row[dst] = text_val(s.get(src))
        stud_rows.append(row)

        status = text_val(s.get("Result Status"))
        if status:
            m = dict(key)
            for src, dst in MS_MAP.items():
                if src in s:
                    m[dst] = text_val(s.get(src))
            m["fetched_at"] = to_iso(s.get("Fetched At"))
            obtained, maximum = rr.split_marks(s.get("Final Total")) if "Final Total" in header else (None, None)
            m["total_obtained"], m["total_max"] = obtained, maximum
            m["percentage"] = round(100.0 * obtained / maximum, 2) if obtained is not None and maximum else None
            extra = {c: text_val(s[c]) for c in header
                     if (c.startswith("Final ") or c.startswith("Marksheet ")) and c not in mapped
                     and text_val(s.get(c)) is not None}
            m["extra"] = extra or None
            if raw and roll in raw:
                m["raw_text"] = raw[roll]
            ms_rows.append(m)

        per_paper = {}
        for col, (paper, comp) in paper_cols.items():
            v = text_val(s.get(col))
            if v is not None:
                per_paper.setdefault(paper, {})[comp] = v
        for paper, comps in per_paper.items():
            total = comps.get("Total")
            pap_rows.append({**key, "paper": paper, "components": comps, "total": total,
                             "total_num": rr.num(total) if total else None})
    return stud_rows, ms_rows, pap_rows


def read_raw(path):
    """roll -> marksheet text, from the 'Raw' sheet of a class workbook (only used with --raw)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    if "Raw" in wb.sheetnames:
        for r in wb["Raw"].iter_rows(min_row=2, values_only=True):
            if r and r[0] is not None:
                out[str(r[0])] = r[1]
    wb.close()
    return out


def class_rows(docs, only_session=None, only_code=None, with_raw=False):
    """Yield (session 'JUN-2026', code, student rows, marksheet rows, paper rows) for every class file."""
    only = only_session.replace("-", "_").upper() if only_session else None
    for sess, code, f in rr.class_files(docs):
        if only and sess.upper() != only:
            continue
        if only_code and code.upper() != only_code.upper():
            continue
        header, students = rr.read_students(f)
        if not students:
            continue
        raw = read_raw(f) if with_raw else None
        st, ms, pp = build_rows(sess.replace("_", "-"), code, header, students, raw)
        yield sess.replace("_", "-"), code, st, ms, pp


# ----------------------------------------------------------------------------
# Sitting-plan PDF -> row
# ----------------------------------------------------------------------------
def parse_sitting_text(text):
    """Read what a sitting-plan PDF says: centre, class line, paper and the roll numbers."""
    flat = "\n".join(l.strip() for l in text.splitlines() if l.strip())
    out = {}
    m = re.search(r"Exam\s*Cent(?:er|re)\s*-\s*(.+)", flat)
    out["exam_centre"] = m.group(1).strip() if m else None
    m = re.search(r"^([^\n/]+(?:/[^\n/]+){3,})$", flat, re.M)
    out["class_line"] = m.group(1).strip() if m else None
    m = re.search(r"Paper\s*Code\s*:\s*([A-Za-z0-9]+)", flat)
    out["paper_code"] = m.group(1) if m else None
    m = re.search(r"Paper\s*Name\s*:\s*(.+?)(?:\s+Date\s*:|\n|$)", flat)
    out["paper_name"] = m.group(1).strip() if m else None
    m = re.search(r"\bTotal\s*:\s*(\d+)", flat)
    out["total_students"] = int(m.group(1)) if m else None

    start, end = re.search(r"Phone\s*No\s*:?", flat), re.search(r"\bPresent\s*:", flat)
    body = flat
    if start:
        body = flat[start.end():end.start() if end and end.start() > start.end() else len(flat)]
    seen, rolls = set(), []
    for r in ROLL_TOKEN.findall(body):
        if r not in seen:
            seen.add(r)
            rolls.append(r)
    out["rolls"], out["roll_count"] = rolls, len(rolls)
    return out


def sitting_row(session, tag, file_key, path):
    info = parse_sitting_text(rr_pdf_text(path))
    parts = (info["class_line"] or "").split("/")
    stem = Path(file_key).stem.split("_")
    row = {
        "session": session, "tag": tag, "file_key": file_key,
        "exam_centre": info["exam_centre"], "class_line": info["class_line"],
        "branch": parts[0].strip() if len(parts) > 3 else "_".join(stem[3:]) or None,
        "sem": parts[1].strip() if len(parts) > 3 else (stem[0] if stem else None),
        "regular_private": parts[2].strip() if len(parts) > 3 else (stem[1] if len(stem) > 1 else None),
        "status": parts[3].strip() if len(parts) > 3 else (stem[2] if len(stem) > 2 else None),
        "paper_code": info["paper_code"], "paper_name": info["paper_name"],
        "total_students": info["total_students"], "roll_count": info["roll_count"], "rolls": info["rolls"],
    }
    return row


def rr_pdf_text(path):
    import pdfplumber
    out = []
    with pdfplumber.open(path) as pdf:
        for pg in pdf.pages:
            out.append(pg.extract_text() or "")
    return "\n".join(out)


def sitting_files(docs, only_session=None):
    """(session 'JUN-2026', tag, 'folder/file.pdf', path) for every sitting-plan PDF below the session folders."""
    only = only_session.upper() if only_session else None
    docs = Path(docs)
    if not docs.is_dir():
        return
    for d in sorted(docs.iterdir()):
        m = SESSION_DIR.fullmatch(d.name)
        if not (d.is_dir() and m):
            continue
        session, tag = f"{m.group(1).upper()}-{m.group(2)}", (m.group(3) or "")
        if only and session != only:
            continue
        for sub in sorted(d.iterdir()):
            if sub.is_dir() and not sub.name.endswith(("_ATTESTATION", "_RESULTS")):
                for f in sorted(sub.glob("*.pdf")):
                    yield session, tag, f"{sub.name}/{f.name}", f


# ----------------------------------------------------------------------------
# Attestation-sheet PDF -> rows (every course; one PDF has many students)
# ----------------------------------------------------------------------------
ATT_ROLL = re.compile(r"Roll\s*No\.?\s*:?\s*(\d{5,})", re.I)
ATT_LABEL = re.compile(
    r"(?<![A-Za-z])(Enrollment\s*No\.?|Session|Regular\s*/\s*Backlog|Exam\s*Name|Exam\s*Cent(?:re|er)|College\s*N\w*"
    r"|Father['’]?s\s*Name|Mother['’]?s\s*Name|Gender|Category|Address|Name)\s*:?", re.I)
ATT_PAPER = re.compile(r"([^\[\]]+?)\s*\[([A-Za-z]+\d+[A-Za-z0-9]*)\]")
ATT_WS = re.compile(r"\s+")


def att_label(raw):
    k = re.sub(r"[^a-z/]", "", raw.lower())
    for start, name in (("enrollment", "Enrollment No"), ("session", "Session"), ("regular", "Regular/Backlog"),
                        ("examcent", "Exam Centre"), ("college", "College Name"), ("father", "Father's Name"),
                        ("mother", "Mother's Name")):
        if k.startswith(start):
            return name
    return {"examname": "Exam Name", "gender": "Gender", "category": "Category", "address": "Address",
            "name": "Name"}.get(k, raw)


def parse_attestation_text(text):
    """One dict per student in the text of an attestation PDF."""
    marks = list(ATT_ROLL.finditer(text))
    out, seen = [], set()
    for i, m in enumerate(marks):
        roll = m.group(1)
        if roll in seen:
            continue
        seen.add(roll)
        end = marks[i + 1].start() if i + 1 < len(marks) else len(text)
        chunk = ATT_WS.sub(" ", text[m.end():end]).strip()
        parts = re.split(r"Paper\s*Description", chunk, maxsplit=1, flags=re.I)
        head, tail = parts[0], (parts[1] if len(parts) > 1 else "")

        rec = {"Roll No": roll}
        found = list(ATT_LABEL.finditer(head))
        for j, lm in enumerate(found):
            nxt = found[j + 1].start() if j + 1 < len(found) else len(head)
            rec.setdefault(att_label(lm.group(1)), head[lm.end():nxt].strip(" :|-"))
        if rec.get("Enrollment No"):
            rec["Enrollment No"] = rec["Enrollment No"].split()[0]
        rec["Address"] = (rec.get("Address") or "")[:300] or None

        tail = re.sub(r"(?<![A-Za-z])(Date|Signature|Answer\s*Sheet\s*No\.?|Invigilator['’]?s?\s*(?:Name|Sign))(?![A-Za-z])",
                      " ", tail, flags=re.I)
        rec["Papers"] = "; ".join(f"{ATT_WS.sub(' ', n).strip(' -:.')} [{c}]" for n, c in ATT_PAPER.findall(tail))
        out.append(rec)
    return out


def exam_code_for(course, sem, status):
    """The result-site exam code of a known course (LLB sem 3 ATKT -> ALL323), else None."""
    try:
        cmap = {k.lower(): v for k, v in getattr(rr.cfg, "COURSE_MAP", {"llb": ("LL", "3")}).items()}
        atkt = [x.upper() for x in getattr(rr.cfg, "ATKT_STATUSES", ["ATKT"])]
        main = [x.upper() for x in getattr(rr.cfg, "MAIN_STATUSES", ["REGULAR", "EXR"])]
    except Exception:
        return None
    m = re.fullmatch(r"(\d+)\s*SEM", str(sem or "").strip().upper())
    st = str(status or "").upper()
    if not m or str(course or "").lower() not in cmap or st not in atkt + main:
        return None
    letters, duration = cmap[str(course).lower()]
    n = int(m.group(1))
    return ("A" if st in atkt else "") + letters + duration + str((n + 1) // 2) + str(n)


def attestation_rows(session, tag, file_key, path):
    """Student rows (attestation_students) of one attestation PDF."""
    stem = Path(file_key).stem.split("_")
    folder = Path(file_key).parent.name
    sem_f, reg_f, status_f = (stem + [None, None, None])[:3]
    rest = "_".join(stem[3:]) or None
    rows = []
    for s in parse_attestation_text(rr_pdf_text(path)):
        exam = [x.strip() for x in (s.get("Exam Name") or "").split(" - ")]
        course = exam[0] if exam and exam[0] else None
        sem = exam[-1] if len(exam) >= 3 else sem_f
        branch = exam[1] if len(exam) >= 3 else rest
        rows.append({
            "session": session, "tag": tag, "file_key": file_key, "roll_no": s["Roll No"],
            "folder": folder or None, "course": course, "branch": branch, "sem": sem,
            "regular_private": reg_f, "status": status_f, "exam_code": exam_code_for(course, sem, status_f),
            "enrollment_no": s.get("Enrollment No") or None, "regular_backlog": s.get("Regular/Backlog") or None,
            "exam_name": s.get("Exam Name") or None, "exam_centre": s.get("Exam Centre") or None,
            "college_name": s.get("College Name") or None, "name": s.get("Name") or None,
            "father_name": s.get("Father's Name") or None, "mother_name": s.get("Mother's Name") or None,
            "gender": s.get("Gender") or None, "category": s.get("Category") or None,
            "address": s.get("Address"), "papers": s.get("Papers") or None,
        })
    return rows


def attestation_files(docs, only_session=None):
    """(session 'JUN-2026', tag, 'sub/file.pdf', path) for every PDF below <SESSION>_ATTESTATION folders."""
    only = only_session.upper() if only_session else None
    docs = Path(docs)
    if not docs.is_dir():
        return
    for d in sorted(docs.iterdir()):
        m = SESSION_DIR.fullmatch(d.name)
        att = d / f"{d.name}_ATTESTATION"
        if not (d.is_dir() and m and att.is_dir()):
            continue
        session, tag = f"{m.group(1).upper()}-{m.group(2)}", (m.group(3) or "")
        if only and session != only:
            continue
        for f in sorted(att.rglob("*.pdf")):
            yield session, tag, f.relative_to(att).as_posix(), f


# ----------------------------------------------------------------------------
# Sync
# ----------------------------------------------------------------------------
def load_state():
    try:
        return json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def save_state(state):
    STATE_FILE.write_text(json.dumps(state, indent=1), encoding="utf-8")


def sync_results(client, docs, only_session=None, only_code=None, with_raw=False, log=print):
    totals = {"students": 0, "marksheets": 0, "marksheet_papers": 0, "classes": 0}
    for session, code, st, ms, pp in class_rows(docs, only_session, only_code, with_raw):
        key = "session,exam_code,roll_no"
        totals["students"] += db.upsert(client, "students", st, key)
        totals["marksheets"] += db.upsert(client, "marksheets", ms, key) if ms else 0
        totals["marksheet_papers"] += db.upsert(client, "marksheet_papers", pp, key + ",paper") if pp else 0
        totals["classes"] += 1
        log(f"   {session} {code}: {len(st)} students, {len(ms)} marksheets, {len(pp)} paper rows")
    if not totals["classes"]:
        log(f"   No class Excel files found under {docs} (<SESSION>\\<SESSION>_RESULTS\\*.xlsx).")
    return totals


def sync_sitting(client, docs, only_session=None, log=print, force=False):
    state = load_state()
    done, skipped, rows = 0, 0, []
    for session, tag, key, f in sitting_files(docs, only_session):
        sig = [f.stat().st_size, int(f.stat().st_mtime)]
        skey = f"{session}|{tag}|{key}"
        if not force and state.get(skey) == sig:
            skipped += 1
            continue
        try:
            row = sitting_row(session, tag, key, f)
        except Exception as e:
            log(f"   {key}: not read ({type(e).__name__}: {str(e)[:80]})")
            continue
        rows.append(row)
        state[skey] = sig
        if len(rows) >= 100:                       # save in batches
            db.upsert(client, "sitting_plans", rows, "session,tag,file_key")
            done += len(rows)
            rows = []
            save_state(state)
            log(f"   sitting plans saved so far: {done}")
    if rows:
        db.upsert(client, "sitting_plans", rows, "session,tag,file_key")
        done += len(rows)
    save_state(state)
    log(f"   sitting plans: {done} new / changed, {skipped} unchanged")
    return {"sitting_plans": done, "unchanged": skipped}


def sync_attestation(client, docs, only_session=None, log=print, force=False):
    """Every attestation PDF (all courses, normal and NEP) -> attestation_students. Unchanged PDFs are skipped."""
    state = load_state()
    students, files, skipped, bad = 0, 0, 0, 0
    pending, sigs = [], {}

    def flush():
        nonlocal students
        if pending:
            students += db.upsert(client, "attestation_students", pending, "session,tag,file_key,roll_no")
            pending.clear()
            state.update(sigs)
            sigs.clear()
            save_state(state)

    for session, tag, key, f in attestation_files(docs, only_session):
        sig = [f.stat().st_size, int(f.stat().st_mtime)]
        skey = f"att|{session}|{tag}|{key}"
        if not force and state.get(skey) == sig:
            skipped += 1
            continue
        try:
            rows = attestation_rows(session, tag, key, f)
        except Exception as e:
            bad += 1
            log(f"   {key}: not read ({type(e).__name__}: {str(e)[:80]})")
            continue
        if not rows:
            bad += 1
            log(f"   {key}: no students found in this PDF")
            continue
        pending.extend(rows)
        sigs[skey] = sig
        files += 1
        if len(pending) >= 300:
            flush()
            log(f"   attestation students saved so far: {students}")
    flush()
    log(f"   attestation: {files} PDF(s) read, {students} student rows saved, {skipped} unchanged"
        + (f", {bad} not readable" if bad else ""))
    return {"attestation_pdfs": files, "attestation_students": students, "attestation_unchanged": skipped}


def sync_all(client, docs, log=print, with_raw=False):
    log("Attestation PDFs:")
    c = sync_attestation(client, docs, log=log)
    log("Results (attestation + marksheet data):")
    a = sync_results(client, docs, with_raw=with_raw, log=log)
    log("Sitting plans:")
    b = sync_sitting(client, docs, log=log)
    return {**c, **a, **b}


def main():
    a = sys.argv[1:]
    mode = a[0].lower() if a else "all"
    docs = Path(rr.DOCS_ROOT) if rr.DOCS_ROOT else rr.documents_dir()
    client = db.connect()
    if mode == "counts":
        for t in db.TABLES:
            print(f"   {t:18} {db.count(client, t)}")
        return
    with_raw = "--raw" in a
    a = [x for x in a if not x.startswith("--")]
    if mode == "results":
        print(sync_results(client, docs, a[1] if len(a) > 1 else None, a[2] if len(a) > 2 else None, with_raw))
    elif mode == "attestation":
        print(sync_attestation(client, docs, a[1] if len(a) > 1 else None, force="--force" in sys.argv))
    elif mode == "sitting":
        print(sync_sitting(client, docs, a[1] if len(a) > 1 else None, force="--force" in sys.argv))
    else:
        print(sync_all(client, docs, with_raw=with_raw))


if __name__ == "__main__":
    main()
