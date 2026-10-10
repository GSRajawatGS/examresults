"""
result_report.py - result report + statistics for every class, made from the Excel files that
result_scraper.py saved (no browser, no internet).

  python result_report.py                 every session found in Documents
  python result_report.py JUN-2023        only that session
  python result_report.py JUN-2023 LL336  only that class (exam code)

For each class  <DOCS>\\JUN_2023\\JUN_2023_RESULTS\\JUN_2023_LL336.xlsx  it writes
  <DOCS>\\JUN_2023\\JUN_2023_RESULTS\\REPORTS\\JUN_2023_LL336_REPORT.xlsx   with the sheets
      Result        one row per student: rank, marks, percentage, division, result (print-ready)
      Statistics    overview, result-wise, division-wise, marks distribution, gender / category /
                    status wise, marks summary, toppers - with charts
      Subjects      paper-wise appeared / absent / highest / lowest / average (and pass % when
                    PAPER_PASS_MARK is set)
      Not passed    follow-up list of students who did not pass, with the papers not cleared
      No result     students of the attestation list whose result was not found / not fetched yet
      Checks        differences between the attestation sheet and the marksheet (name, father, ...)
and one file for all classes together:
  <DOCS>\\RESULT_STATISTICS.xlsx
      All classes   one row per class of every session
      By session    main / ATKT / all pass % of every session (with chart)
      Trend         pass % of the same class (e.g. LLB sem 3 main) across sessions
      Toppers       the first three of every class

It is also run by "python result_scraper.py report", and by "auto" after new results arrive.
Settings (all optional, in result_jobs.py): PASS_WORDS, NOT_PASS_WORDS, TOP_N, DIVISIONS, BANDS,
PAPER_PASS_MARK, ABSENT_WORDS - see the README block at the end of this file.
"""
import re
import statistics
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

try:
    import result_jobs as cfg
except Exception:                       # works without result_jobs.py too
    cfg = None

DOCS_ROOT = getattr(cfg, "DOCS_ROOT", None)
# A "Final Result" value counts as PASSED if it contains one of these words and none of the NOT_PASS words
PASS_WORDS = tuple(w.upper() for w in getattr(cfg, "PASS_WORDS", ("PASS",)))
NOT_PASS_WORDS = tuple(w.upper() for w in getattr(
    cfg, "NOT_PASS_WORDS", ("FAIL", "ATKT", "ABSENT", "WITHHELD", "CANCEL", "UFM")))
TOP_N = int(getattr(cfg, "TOP_N", 10))
# Division of a PASSED student from the percentage: (minimum %, name), highest first.
DIVISIONS = tuple(sorted(((float(a), str(b)) for a, b in getattr(
    cfg, "DIVISIONS", ((60, "First"), (48, "Second"), (0, "Third")))), reverse=True))
# Edges of the marks-distribution bands, in percent: (40, 50, 60, 70) -> <40, 40-50, 50-60, 60-70, 70+
BANDS = tuple(sorted(float(x) for x in getattr(cfg, "BANDS", (40, 50, 60, 70))))
# Pass mark of one paper (same unit as the paper totals, e.g. 36 of 100). None = paper pass/fail is not computed.
PAPER_PASS_MARK = getattr(cfg, "PAPER_PASS_MARK", None)
ABSENT_WORDS = tuple(w.upper() for w in getattr(cfg, "ABSENT_WORDS", ("AB", "ABS", "ABSENT")))

MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
FONT = "Arial"
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="305496")
SECTION_FONT = Font(name=FONT, bold=True, size=12, color="305496")
TITLE_FONT = Font(name=FONT, bold=True, size=14)
NOTE_FONT = Font(name=FONT, italic=True, size=9, color="595959")
NUM_RE = re.compile(r"\d+(?:\.\d+)?")


def documents_dir():
    try:
        import ctypes
        buf = ctypes.create_unicode_buffer(260)
        ctypes.windll.shell32.SHGetFolderPathW(None, 5, None, 0, buf)
        return Path(buf.value)
    except Exception:
        return Path.home() / "Documents"


# ----------------------------------------------------------------------------
# Reading one class workbook
# ----------------------------------------------------------------------------
def num(value):
    """First number inside a cell ('456', '456/700', '45.5'); None if there is none (AB, --, empty)."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    m = NUM_RE.search(str(value))
    return float(m.group(0)) if m else None


def split_marks(value):
    """'253/500' -> (253.0, 500.0); '253' -> (253.0, None); 'AB' -> (None, None)."""
    if value is None:
        return None, None
    if isinstance(value, (int, float)):
        return float(value), None
    text = str(value)
    nums = NUM_RE.findall(text)
    if not nums:
        return None, None
    return float(nums[0]), (float(nums[1]) if "/" in text and len(nums) > 1 else None)


def read_students(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Students"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    header = [str(h) if h is not None else "" for h in rows[0]]
    students = []
    for r in rows[1:]:
        if r and r[0] is not None:
            students.append({h: v for h, v in zip(header, r) if h})
    return header, students


def find_columns(header):
    """Which columns hold the final result / final total / paper totals (names come from the result site)."""
    low = {h: h.lower() for h in header}
    final_cols = [h for h in header if low[h].startswith("final ")]
    result_col = next((h for h in final_cols if "result" in low[h]), None)
    total_col = next((h for h in final_cols if "total" in low[h]), None)
    if total_col is None:                                    # fall back to a plain 'Total' column
        total_col = next((h for h in header if low[h] in ("total", "grand total")), None)
    paper_cols = [h for h in header if " | " in h and h.split(" | ")[-1].strip() == "Total"]
    other_final = [h for h in final_cols if h not in (result_col, total_col)]
    return result_col, total_col, paper_cols, other_final


def is_pass(text):
    t = str(text or "").upper()
    return any(w in t for w in PASS_WORDS) and not any(w in t for w in NOT_PASS_WORDS)


def clean(text, default="(blank)"):
    t = " ".join(str(text).split()) if text not in (None, "") else ""
    return t.upper() if t else default


def exam_type(code, students):
    """ATKT or Main: the 'Class' column written by result_scraper says it; else an A in front of the code."""
    for s in students:
        label = str(s.get("Class") or "")
        if label:
            return "ATKT" if "ATKT" in label.upper() else "Main"
    return "ATKT" if code.upper().startswith("A") else "Main"


def pct(a, b):
    return round(100.0 * a / b, 2) if b else 0.0


def division_of(percent):
    for threshold, name in DIVISIONS:
        if percent >= threshold:
            return name
    return DIVISIONS[-1][1] if DIVISIONS else ""


def band_labels():
    edges = list(BANDS)
    if not edges:
        return ["All"]
    labels = [f"Below {edges[0]:g}%"]
    labels += [f"{a:g}% to below {b:g}%" for a, b in zip(edges, edges[1:])]
    labels.append(f"{edges[-1]:g}% and above")
    return labels


def band_index(percent):
    return sum(1 for e in BANDS if percent >= e)


def is_absent(value):
    return str(value or "").strip().upper() in ABSENT_WORDS


def norm_key(text):
    return re.sub(r"[^A-Z0-9]", "", str(text or "").upper())


# ----------------------------------------------------------------------------
# Statistics
# ----------------------------------------------------------------------------
def compute(header, students):
    result_col, total_col, paper_cols, other_final = find_columns(header)
    ok = [s for s in students if str(s.get("Result Status") or "") == "OK"]
    not_found = [s for s in students if str(s.get("Result Status") or "") == "NOT FOUND"]
    no_result = [s for s in students if str(s.get("Result Status") or "") != "OK"]
    pending = len(students) - len(ok) - len(not_found)

    for s in ok:
        obtained, maximum = split_marks(s.get(total_col)) if total_col else (None, None)
        s["_total"], s["_max"] = obtained, maximum
        s["_pct"] = round(100.0 * obtained / maximum, 2) if obtained is not None and maximum else None
        res = clean(s.get(result_col)) if result_col else "(no result column)"
        s["_result"] = res
        s["_pass"] = is_pass(res)
        s["_div"] = division_of(s["_pct"]) if s["_pass"] and s["_pct"] is not None else ""
        failed = []
        if PAPER_PASS_MARK is not None:
            for col in paper_cols:
                v = num(s.get(col))
                if v is not None and not is_absent(s.get(col)) and v < float(PAPER_PASS_MARK):
                    failed.append(col.rsplit(" | ", 1)[0])
        s["_failed_papers"] = failed

    ranked = sorted((s for s in ok if s["_total"] is not None), key=lambda s: -s["_total"])
    rank, last, pos = {}, None, 0
    for i, s in enumerate(ranked, 1):                         # same marks = same rank
        if s["_total"] != last:
            pos, last = i, s["_total"]
        rank[id(s)] = pos

    def group(col, default="(unknown)"):
        out = {}
        for s in students:
            key = clean(s.get(col), default)
            g = out.setdefault(key, {"listed": 0, "appeared": 0, "passed": 0, "pcts": []})
            g["listed"] += 1
            if str(s.get("Result Status") or "") == "OK":
                g["appeared"] += 1
                g["passed"] += 1 if s["_pass"] else 0
                if s["_pct"] is not None:
                    g["pcts"].append(s["_pct"])
        return out

    groups = []
    for title, col in (("Gender-wise", "Gender"), ("Category-wise", "Category"),
                       ("Regular / Backlog", "Regular/Backlog"), ("Status (attestation list)", "Att Status"),
                       ("Regular / Private (attestation list)", "Att Regular/Private"),
                       ("Branch", "Att Branch"), ("Exam code (chance)", "Result Exam Code")):
        if col in header:
            g = group(col)
            if len(g) >= 2 or col == "Gender":                # a one-row table says nothing
                groups.append((title, g))

    results = {}
    for s in ok:
        results[s["_result"]] = results.get(s["_result"], 0) + 1

    other = {}
    for col in other_final:
        counts = {}
        for s in ok:
            key = clean(s.get(col))
            counts[key] = counts.get(key, 0) + 1
        if len(counts) <= 12:                                 # a category-like column (division, grade ...)
            other[col] = counts

    divisions = {name: 0 for _, name in DIVISIONS}
    for s in ok:
        if s["_div"]:
            divisions[s["_div"]] += 1
    no_percent_pass = sum(1 for s in ok if s["_pass"] and s["_pct"] is None)

    bands = [0] * len(band_labels())
    for s in ok:
        if s["_pct"] is not None:
            bands[band_index(s["_pct"])] += 1

    totals = [s["_total"] for s in ok if s["_total"] is not None]
    pcts = [s["_pct"] for s in ok if s["_pct"] is not None]
    max_marks = None
    if any(s["_max"] for s in ok):
        vals = [s["_max"] for s in ok if s["_max"]]
        max_marks = max(set(vals), key=vals.count)

    subjects = []
    for col in paper_cols:
        raw = [s.get(col) for s in ok]
        entered = [v for v in raw if v not in (None, "")]
        absent = [v for v in entered if is_absent(v)]
        vals_n = [num(v) for v in entered if not is_absent(v) and num(v) is not None]
        sj = {"paper": col.rsplit(" | ", 1)[0], "appeared": len(entered) - len(absent),
              "absent": len(absent), "numeric": len(vals_n),
              "highest": max(vals_n) if vals_n else None, "lowest": min(vals_n) if vals_n else None,
              "average": round(statistics.mean(vals_n), 2) if vals_n else None,
              "failed": None, "pass_pct": None}
        if PAPER_PASS_MARK is not None and vals_n:
            sj["failed"] = sum(1 for v in vals_n if v < float(PAPER_PASS_MARK))
            sj["pass_pct"] = pct(len(vals_n) - sj["failed"], len(vals_n))
        subjects.append(sj)

    checks = []
    pairs = (("Name", "Marksheet Name"), ("Father's Name", "Marksheet Father's Name"),
             ("Mother's Name", "Marksheet Mother's Name"),
             ("Enrollment No", "Marksheet Enrollment Number"), ("Roll No", "Marksheet Roll Number"))
    for att, ms in pairs:
        if att in header and ms in header:
            for s in ok:
                a, m = s.get(att), s.get(ms)
                if a not in (None, "") and m not in (None, "") and norm_key(a) != norm_key(m):
                    checks.append((s.get("Roll No"), att, a, m))

    return {
        "result_col": result_col, "total_col": total_col, "listed": len(students), "ok": ok,
        "appeared": len(ok), "not_found": len(not_found), "pending": pending, "no_result": no_result,
        "passed": sum(1 for s in ok if s["_pass"]), "results": results, "groups": groups, "other": other,
        "totals": totals, "pcts": pcts, "max_marks": max_marks, "ranked": ranked, "rank": rank,
        "subjects": subjects, "paper_cols": paper_cols, "divisions": divisions,
        "no_percent_pass": no_percent_pass, "bands": bands, "checks": checks,
    }


# ----------------------------------------------------------------------------
# Writing helpers
# ----------------------------------------------------------------------------
def head(ws, row, values, col=1):
    for c, v in enumerate(values, col):
        cell = ws.cell(row, c, v)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, maxw=45, minw=10):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None and c.row > 0), default=8)
        ws.column_dimensions[letter].width = min(max(minw, longest + 2), maxw)


def use_font(ws):
    """Arial everywhere, keeping bold / size / colour."""
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != FONT:
                f = copy(c.font)
                c.font = Font(name=FONT, size=f.size if f.size and f.size != 11 else 10, bold=f.bold,
                              italic=f.italic, color=f.color)


def print_setup(ws, title_rows=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if title_rows:
        ws.print_title_rows = title_rows


class Sheet:
    """Writes titles, sections and tables down a sheet and remembers where it is."""

    def __init__(self, ws):
        self.ws, self.row = ws, 1

    def title(self, text):
        self.ws.cell(self.row, 1, text).font = TITLE_FONT
        self.row += 1

    def section(self, text):
        self.row += 1
        self.ws.cell(self.row, 1, text).font = SECTION_FONT
        self.row += 1

    def note(self, text):
        self.ws.cell(self.row, 1, text).font = NOTE_FONT
        self.row += 1

    def kv(self, pairs, fmt=None):
        for k, v in pairs:
            self.ws.cell(self.row, 1, k)
            c = self.ws.cell(self.row, 2, v)
            if fmt and k in fmt:
                c.number_format = fmt[k]
            self.row += 1

    def table(self, headers, rows, formats=None):
        """Returns (header_row, last_row). formats: {column index (1-based): number format}."""
        top = self.row
        head(self.ws, top, headers)
        self.row += 1
        for r in rows:
            for c, v in enumerate(r, 1):
                cell = self.ws.cell(self.row, c, v)
                if formats and c in formats:
                    cell.number_format = formats[c]
            self.row += 1
        return top, self.row - 1


def add_chart(ws, kind, title, header_row, last_row, cat_col, data_cols, anchor, y_title=None, legend=False,
              width=14, height=7.5, flow=True):
    """A chart over a table made by Sheet.table (header row included, so series get their names).
    flow=True: the chart goes at or below the wanted row of the column in `anchor`, never on top of
    an earlier chart."""
    if last_row <= header_row:
        return
    if flow:
        m = re.fullmatch(r"([A-Z]+)(\d+)", anchor)
        row = max(int(m.group(2)), getattr(ws, "_next_chart_row", 2))
        ws._next_chart_row = row + int(height * 2) + 2
        anchor = f"{m.group(1)}{row}"
    if kind == "pie":
        ch = PieChart()
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
        ch.dataLabels.showVal = False
        ch.dataLabels.showCatName = False
        ch.dataLabels.showSerName = False
        ch.dataLabels.showLeaderLines = False
    elif kind == "line":
        ch = LineChart()
        ch.y_axis.delete = False
        ch.x_axis.delete = False
    else:
        ch = BarChart()
        ch.type = "col"
        ch.y_axis.delete = False
        ch.x_axis.delete = False
        ch.y_axis.majorGridlines = None
    ch.title = title
    for dc in data_cols:
        ch.add_data(Reference(ws, min_col=dc, min_row=header_row, max_row=last_row), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=header_row + 1, max_row=last_row))
    if y_title and kind != "pie":
        ch.y_axis.title = y_title
        if y_title.startswith("Pass"):                    # percentages always on a 0-100 scale
            ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
    if not legend and kind != "pie":
        ch.legend = None
    ch.width, ch.height = width, height
    ws.add_chart(ch, anchor)


# ----------------------------------------------------------------------------
# One class: the report workbook
# ----------------------------------------------------------------------------
def write_class_report(path, session, code, st, header, students, source_name=""):
    kind = exam_type(code, students)
    wb = Workbook()

    # ------------------------------------------------------------ Result
    ws = wb.active
    ws.title = "Result"
    info = [c for c in ["Roll No", "Name", "Father's Name", "Enrollment No", "Gender", "Category",
                        "Regular/Backlog", "Result Exam Code"] if c in header]
    extra = ["Percentage", "Division"]
    cols = (["Rank"] + info + st["paper_cols"] + ([st["total_col"]] if st["total_col"] else []) + extra
            + ([st["result_col"]] if st["result_col"] and st["result_col"] != st["total_col"] else [])
            + (["Papers not cleared"] if PAPER_PASS_MARK is not None else [])
            + ["Result Status", "Result Note"])
    titles = [c.rsplit(" | ", 1)[0] + " (total)" if c in st["paper_cols"] else c for c in cols]
    head(ws, 1, titles)
    ordered = sorted(students, key=lambda s: (str(s.get("Result Status")) != "OK",
                                              st["rank"].get(id(s), 10 ** 9), str(s.get("Roll No"))))
    for r, s in enumerate(ordered, 2):
        for c, name in enumerate(cols, 1):
            if name == "Rank":
                v = st["rank"].get(id(s))
            elif name == "Percentage":
                v = s.get("_pct")
            elif name == "Division":
                v = s.get("_div") or None
            elif name == "Papers not cleared":
                v = ", ".join(s.get("_failed_papers", [])) or None
            else:
                v = s.get(name)
                if name in st["paper_cols"] or name == st["total_col"]:
                    n = num(v)
                    v = n if n is not None and str(v).strip().replace(".", "", 1).isdigit() else v
            cell = ws.cell(r, c, v)
            if name == "Percentage":
                cell.number_format = "0.00"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    widths(ws)
    print_setup(ws, "1:1")

    # ------------------------------------------------------------ Statistics
    ws = wb.create_sheet("Statistics")
    sh = Sheet(ws)
    sh.title(f"{session}  |  {code}  |  {kind} exam")
    sh.note("Percentages are of the students whose result was found, unless a column says otherwise.")

    sh.section("Overview")
    sh.kv([("Students in attestation list", st["listed"]),
           ("Result found (appeared)", st["appeared"]),
           ("Result not found", st["not_found"]),
           ("Not looked up yet / error", st["pending"]),
           ("Passed", st["passed"]),
           ("Not passed", st["appeared"] - st["passed"]),
           ("Pass % (of result found)", pct(st["passed"], st["appeared"])),
           ("Pass % (of attestation list)", pct(st["passed"], st["listed"]))])
    pie_top = sh.row + 1
    sh.row += 1
    pass_hdr, pass_last = sh.table(["Outcome", "Students"],
                                   [("Passed", st["passed"]), ("Not passed", st["appeared"] - st["passed"])])
    add_chart(ws, "pie", "Passed / not passed", pass_hdr, pass_last, 1, [2], f"H{max(pie_top - 6, 2)}",
              width=11, height=6.5)

    sh.section("Result-wise (as printed on the marksheet)")
    sh.table(["Result", "Students", "% of appeared"],
             [(k, v, pct(v, st["appeared"])) for k, v in sorted(st["results"].items(), key=lambda x: -x[1])])

    if st["pcts"]:
        sh.section("Division-wise (passed students; worked out from the marks)")
        rows = [(name, n, pct(n, st["appeared"]), pct(n, st["passed"])) for name, n in st["divisions"].items()]
        if st["no_percent_pass"]:
            rows.append(("(percentage not available)", st["no_percent_pass"], pct(st["no_percent_pass"], st["appeared"]),
                         pct(st["no_percent_pass"], st["passed"])))
        d_hdr, d_last = sh.table(["Division", "Students", "% of appeared", "% of passed"], rows)
        add_chart(ws, "bar", "Division-wise", d_hdr, d_last, 1, [2], f"H{d_hdr - 2}", y_title="Students",
                  width=11, height=6.5)

        sh.section("Marks distribution (percentage" + (f" of {st['max_marks']:g} marks" if st["max_marks"] else "") + ")")
        labels = band_labels()
        b_hdr, b_last = sh.table(["Marks band", "Students", "% of appeared"],
                                 [(lab, n, pct(n, st["appeared"])) for lab, n in zip(labels, st["bands"])])
        add_chart(ws, "bar", "Marks distribution", b_hdr, b_last, 1, [2], f"H{b_hdr - 1}", y_title="Students",
                  width=11, height=6.5)

    for title, data in st["groups"]:
        sh.section(title)
        rows = []
        for k, g in sorted(data.items()):
            rows.append((k, g["listed"], g["appeared"], g["passed"], pct(g["passed"], g["appeared"]),
                         round(statistics.mean(g["pcts"]), 2) if g["pcts"] else None))
        first = title.split(" (")[0].replace("-wise", "")
        g_hdr, g_last = sh.table([first, "In list", "Appeared", "Passed", "Pass % (of appeared)", "Average %"], rows)
        if title in ("Gender-wise", "Category-wise") and g_last > g_hdr:
            add_chart(ws, "bar", f"Pass % - {title.lower()}", g_hdr, g_last, 1, [5], f"H{g_hdr - 1}",
                      y_title="Pass %", width=11, height=6.5)

    for col, counts in st["other"].items():
        sh.section(col)
        sh.table(["Value", "Students", "% of appeared"],
                 [(k, v, pct(v, st["appeared"])) for k, v in sorted(counts.items(), key=lambda x: -x[1])])

    if st["totals"]:
        sh.section(f"Marks ({st['total_col']})")
        rows = [("Highest", max(st["totals"])), ("Lowest", min(st["totals"])),
                ("Average", round(statistics.mean(st["totals"]), 2)),
                ("Median", round(statistics.median(st["totals"]), 2))]
        if st["pcts"]:
            rows += [("Highest %", max(st["pcts"])), ("Lowest %", min(st["pcts"])),
                     ("Average %", round(statistics.mean(st["pcts"]), 2))]
        passed_t = [s["_total"] for s in st["ok"] if s["_pass"] and s["_total"] is not None]
        if passed_t:
            rows.append(("Average of passed students", round(statistics.mean(passed_t), 2)))
        sh.kv(rows)

        sh.section(f"Top {TOP_N}")
        top_rows = []
        for s in [x for x in st["ranked"] if st["rank"][id(x)] <= TOP_N][:TOP_N + 5]:
            top_rows.append((st["rank"][id(s)], s.get("Roll No"), s.get("Name"), s["_total"], s["_pct"],
                             s["_div"] or None, s["_result"]))
        sh.table(["Rank", "Roll No", "Name", "Total", "%", "Division", "Result"], top_rows)

    sh.section("How these numbers were worked out")
    sh.note(f"Pass = the marksheet result contains {'/'.join(PASS_WORDS)} and none of {'/'.join(NOT_PASS_WORDS)}.")
    sh.note("Percentage = obtained / maximum of the 'Final Total' (for example 253/500). "
            "Division is worked out from it for PASSED students only: "
            + ", ".join(f"{n} {a:g}% and above" if a else f"{n} below that" for a, n in DIVISIONS)
            + " - it can differ from the university's own division.")
    if PAPER_PASS_MARK is not None:
        sh.note(f"A paper counts as not cleared when its marks are below {PAPER_PASS_MARK}.")
    sh.note(f"Source: {source_name or 'class Excel file'}  |  made {datetime.now():%d-%m-%Y %H:%M}")
    widths(ws, 40)
    ws.column_dimensions["A"].width = 38
    print_setup(ws)

    # ------------------------------------------------------------ Subjects
    ws = wb.create_sheet("Subjects")
    hdrs = ["Paper", "Appeared", "Absent", "Highest", "Lowest", "Average"]
    if PAPER_PASS_MARK is not None:
        hdrs += ["Below pass mark", "Pass %"]
    head(ws, 1, hdrs)
    for r, sj in enumerate(st["subjects"], 2):
        vals = [sj["paper"], sj["appeared"], sj["absent"], sj["highest"], sj["lowest"], sj["average"]]
        if PAPER_PASS_MARK is not None:
            vals += [sj["failed"], sj["pass_pct"]]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
    if not st["subjects"]:
        ws.cell(2, 1, "No paper-wise totals were found in this Excel file.")
    else:
        last = len(st["subjects"]) + 1
        if PAPER_PASS_MARK is not None:
            add_chart(ws, "bar", "Pass % by paper", 1, last, 1, [8], f"A{last + 3}", y_title="Pass %",
                      width=22, height=8, flow=False)
        else:
            add_chart(ws, "bar", "Average marks by paper", 1, last, 1, [6], f"A{last + 3}", y_title="Average",
                      width=22, height=8, flow=False)
    widths(ws, 60)
    print_setup(ws, "1:1")

    # ------------------------------------------------------------ Not passed
    ws = wb.create_sheet("Not passed")
    np_cols = [c for c in ["Roll No", "Name", "Father's Name", "Gender", "Category"] if c in header]
    head(ws, 1, np_cols + ["Result", "Total", "%"] + (["Papers not cleared"] if PAPER_PASS_MARK is not None else []))
    failed_students = sorted((s for s in st["ok"] if not s["_pass"]), key=lambda s: str(s.get("Roll No")))
    for r, s in enumerate(failed_students, 2):
        vals = [s.get(c) for c in np_cols] + [s["_result"], s["_total"], s["_pct"]]
        if PAPER_PASS_MARK is not None:
            vals.append(", ".join(s["_failed_papers"]) or None)
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
    if not failed_students:
        ws.cell(2, 1, "Every student whose result was found has passed.")
    ws.freeze_panes = "A2"
    widths(ws, 50)
    print_setup(ws, "1:1")

    # ------------------------------------------------------------ No result
    ws = wb.create_sheet("No result")
    nr_cols = [c for c in ["Roll No", "Name", "Father's Name", "Enrollment No"] if c in header]
    tail = [c for c in ["Result Status", "Result Note", "Fetched At", "Source File"] if c in header]
    head(ws, 1, nr_cols + tail)
    rows_nr = sorted(st["no_result"], key=lambda s: str(s.get("Roll No")))
    for r, s in enumerate(rows_nr, 2):
        for c, name in enumerate(nr_cols + tail, 1):
            ws.cell(r, c, s.get(name) or ("(not looked up yet)" if name == "Result Status" else None))
    if not rows_nr:
        ws.cell(2, 1, "A result was found for every student of the attestation list.")
    ws.freeze_panes = "A2"
    widths(ws, 50)
    print_setup(ws, "1:1")

    # ------------------------------------------------------------ Checks
    ws = wb.create_sheet("Checks")
    ws.cell(1, 1, "Attestation sheet compared with the marksheet (spelling differences are normal; "
                  "a different person is not).").font = NOTE_FONT
    head(ws, 2, ["Roll No", "Field", "Attestation sheet", "Marksheet"])
    for r, row in enumerate(st["checks"], 3):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    if not st["checks"]:
        ws.cell(3, 1, "No differences found.")
    widths(ws, 50)

    for sheet in wb.worksheets:
        use_font(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ----------------------------------------------------------------------------
# All classes together
# ----------------------------------------------------------------------------
def session_key(name):
    m = re.fullmatch(r"([A-Za-z]{3})_(\d{4})", name)
    return (int(m.group(2)), MONTHS.index(m.group(1).upper()) if m.group(1).upper() in MONTHS else 0) if m else (0, 0)


def class_files(docs):
    """[(session 'JUN_2023', exam code, xlsx path)] for every class Excel file."""
    found = []
    if not docs.is_dir():
        return found
    for d in sorted((x for x in docs.iterdir() if x.is_dir() and re.fullmatch(r"[A-Za-z]{3}_\d{4}", x.name)),
                    key=lambda x: session_key(x.name)):
        res = d / f"{d.name}_RESULTS"
        if not res.is_dir():
            continue
        for f in sorted(res.glob(f"{d.name}_*.xlsx")):
            code = f.stem[len(d.name) + 1:]
            if f.name.startswith("~$") or code.upper().endswith("_REPORT") or not re.fullmatch(r"A?\d?[A-Za-z]+\d+", code):
                continue
            found.append((d.name, code, f))
    return found


def family_label(code, students):
    """'LLB sem 3 main' - the same class in every session, so sessions can be compared."""
    for s in students:
        label = str(s.get("Class") or "")
        if label:
            return re.sub(r"\s*\([^)]*\)\s*$", "", label).strip()
    m = re.fullmatch(r"(A?)([A-Za-z]+?)(\d)(\d)(\d)", code)
    if m:
        return f"{m.group(2).upper()}{m.group(3)} sem {m.group(5)} {'ATKT' if m.group(1) else 'main'}"
    return code


def summary_record(sess, code, students, st):
    avg_pct = round(statistics.mean(st["pcts"]), 2) if st["pcts"] else None
    return {
        "session": sess.replace("_", "-"), "code": code, "kind": exam_type(code, students),
        "family": family_label(code, students), "listed": st["listed"], "appeared": st["appeared"],
        "not_found": st["not_found"], "pending": st["pending"], "passed": st["passed"],
        "highest": max(st["totals"]) if st["totals"] else None,
        "average": round(statistics.mean(st["totals"]), 2) if st["totals"] else None,
        "avg_pct": avg_pct, "divisions": st["divisions"],
        "top": [(st["rank"][id(s)], s.get("Roll No"), s.get("Name"), s["_total"], s["_pct"], s["_div"] or None)
                for s in st["ranked"][:3]],
    }


def write_summary(path, records):
    wb = Workbook()
    div_names = [n for _, n in DIVISIONS]

    # ------------------------------------------------------------ All classes
    ws = wb.active
    ws.title = "All classes"
    cols = ["Session", "Exam code", "Exam", "In list", "Appeared", "Not found", "Pending", "Passed",
            "Not passed", "Pass % (of appeared)", "Highest", "Average"] + ["Average %"] + div_names
    head(ws, 1, cols)
    for r, x in enumerate(records, 2):
        row = [x["session"], x["code"], x["kind"], x["listed"], x["appeared"], x["not_found"], x["pending"],
               x["passed"], x["appeared"] - x["passed"], pct(x["passed"], x["appeared"]), x["highest"], x["average"],
               x["avg_pct"]] + [x["divisions"].get(n, 0) for n in div_names]
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths(ws)
    print_setup(ws, "1:1")

    sessions = sorted({x["session"] for x in records}, key=lambda s: session_key(s.replace("-", "_")))

    # ------------------------------------------------------------ By session
    ws = wb.create_sheet("By session")
    head(ws, 1, ["Session", "Classes", "Main: appeared", "Main: passed", "Main: pass %",
                 "ATKT: appeared", "ATKT: passed", "ATKT: pass %", "All: appeared", "All: passed", "All: pass %"])
    for r, sess in enumerate(sessions, 2):
        rec = [x for x in records if x["session"] == sess]
        main = [x for x in rec if x["kind"] == "Main"]
        atkt = [x for x in rec if x["kind"] == "ATKT"]
        ma, mp = sum(x["appeared"] for x in main), sum(x["passed"] for x in main)
        aa, ap = sum(x["appeared"] for x in atkt), sum(x["passed"] for x in atkt)
        row = [sess, len(rec), ma, mp, pct(mp, ma) if ma else None, aa, ap, pct(ap, aa) if aa else None,
               ma + aa, mp + ap, pct(mp + ap, ma + aa) if ma + aa else None]
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    last = len(sessions) + 1
    if last >= 3:
        ch = LineChart()
        ch.title = "Pass % by session"
        ch.y_axis.title = "Pass %"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 100
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        for dc in (5, 8, 11):
            ch.add_data(Reference(ws, min_col=dc, min_row=1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
        for series in ch.series:
            series.smooth = False                         # straight lines between sessions
        ch.width, ch.height = 22, 9
        ws.add_chart(ch, f"A{last + 3}")
    widths(ws, 24)
    print_setup(ws, "1:1")

    # ------------------------------------------------------------ Trend
    ws = wb.create_sheet("Trend")
    ws.cell(1, 1, "Pass % of the same class in every session (empty = that class had no result in the session)."
            ).font = NOTE_FONT
    head(ws, 2, ["Class"] + sessions + ["Change vs previous (points)"])
    families = sorted({x["family"] for x in records})
    change_col = len(sessions) + 2
    for r, fam in enumerate(families, 3):
        ws.cell(r, 1, fam)
        values = []
        for c, sess in enumerate(sessions, 2):
            recs = [x for x in records if x["family"] == fam and x["session"] == sess]
            a, p_ = sum(x["appeared"] for x in recs), sum(x["passed"] for x in recs)
            if a:
                ws.cell(r, c, pct(p_, a))
                values.append(pct(p_, a))
        if len(values) >= 2:                              # latest session with a result vs the one before
            cell = ws.cell(r, change_col, round(values[-1] - values[-2], 2))
            cell.number_format = "+0.00;-0.00;0.00"
    if families and sessions:
        rng = f"B3:{get_column_letter(len(sessions) + 1)}{len(families) + 2}"
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                                           mid_type="num", mid_value=50, mid_color="FFEB84",
                                                           end_type="num", end_value=100, end_color="63BE7B"))
    ws.freeze_panes = "B3"
    widths(ws, 30)
    print_setup(ws, "2:2")

    # ------------------------------------------------------------ Toppers
    ws = wb.create_sheet("Toppers")
    head(ws, 1, ["Session", "Exam code", "Class", "Rank", "Roll No", "Name", "Total", "%", "Division"])
    r = 2
    for x in records:
        for rank, roll, name, total, p_, div in x["top"]:
            for c, v in enumerate([x["session"], x["code"], x["family"], rank, roll, name, total, p_, div], 1):
                ws.cell(r, c, v)
            r += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths(ws, 40)
    print_setup(ws, "1:1")

    for sheet in wb.worksheets:
        use_font(sheet)
    wb.save(path)


def build_summary(docs, quiet=False):
    """RESULT_STATISTICS.xlsx: every class of every session (always covers all classes)."""
    records = []
    for sess, code, f in class_files(docs):
        try:
            header, students = read_students(f)
            if students:
                records.append(summary_record(sess, code, students, compute(header, students)))
        except Exception as e:
            print(f"   {f.name}: not in the summary ({type(e).__name__}: {e})")
    if not records:
        print(f"   No class Excel files found under {docs} (looking for <SESSION>\\<SESSION>_RESULTS\\*.xlsx).")
        return None
    out = docs / "RESULT_STATISTICS.xlsx"
    try:
        write_summary(out, records)
    except PermissionError:
        print("   RESULT_STATISTICS.xlsx is open in Excel - close it and run again.")
        return None
    if not quiet:
        print(f"\nAll-class statistics ({len(records)} classes) -> {out}")
    return out


def build_reports(docs=None, only_session=None, only_code=None, quiet=False):
    docs = Path(docs) if docs else (Path(DOCS_ROOT) if DOCS_ROOT else documents_dir())
    only_session = only_session.replace("-", "_").upper() if only_session else None
    made = 0
    for sess, code, f in class_files(docs):
        if only_session and sess.upper() != only_session:
            continue
        if only_code and code.upper() != only_code.upper():
            continue
        try:
            header, students = read_students(f)
            if not students:
                if not quiet:
                    print(f"   {f.name}: no students, skipped")
                continue
            st = compute(header, students)
            out = f.parent / "REPORTS" / f"{sess}_{code}_REPORT.xlsx"
            write_class_report(out, sess.replace("_", "-"), code, st, header, students, f.name)
            made += 1
            if not quiet:
                note = "" if st["result_col"] else "   (no 'Final ...Result' column yet - results not fetched?)"
                print(f"   {sess} {code}: {st['appeared']}/{st['listed']} results, {st['passed']} passed "
                      f"({pct(st['passed'], st['appeared'])}%) -> {out}{note}")
        except PermissionError:
            print(f"   {f.name}: report file is open in Excel - close it and run again.")
        except Exception as e:
            print(f"   {f.name}: FAILED ({type(e).__name__}: {e})")
    if not quiet:
        print(f"\n{made} class report(s) written.")
    build_summary(docs, quiet=quiet)          # always refreshed, whichever session/class was asked for
    return made


# README ----------------------------------------------------------------------
# Optional settings in result_jobs.py (the defaults are shown):
#   PASS_WORDS      = ("PASS",)
#   NOT_PASS_WORDS  = ("FAIL", "ATKT", "ABSENT", "WITHHELD", "CANCEL", "UFM")
#   TOP_N           = 10
#   DIVISIONS       = ((60, "First"), (48, "Second"), (0, "Third"))   # minimum %, name - passed students only
#   BANDS           = (40, 50, 60, 70)                                # marks-distribution edges in %
#   PAPER_PASS_MARK = None      # e.g. 36 -> paper-wise pass %, "Papers not cleared" columns and list
#   ABSENT_WORDS    = ("AB", "ABS", "ABSENT")

if __name__ == "__main__":
    a = sys.argv[1:]
    build_reports(only_session=a[0] if a else None, only_code=a[1] if len(a) > 1 else None)
